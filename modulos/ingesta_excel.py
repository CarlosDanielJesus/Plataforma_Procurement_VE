import re
import unicodedata

import pandas as pd
import streamlit as st

from modulos.imagenes import normalizar_url_imagen

COLUMNAS_REQUERIDAS = ['SKU', 'MATERIAL', 'PRECIO_USD', 'STOCK']

# Alias aceptados (ya normalizados: mayúsculas, sin tildes, sin espacios ni símbolos)
ALIAS = {
    'SKU': {'SKU', 'CODIGO', 'COD'},
    'MATERIAL': {'MATERIAL', 'DESCRIPCION', 'PRODUCTO', 'NOMBRE'},
    'PRECIO_USD': {'PRECIOUSD', 'PRECIO', 'PRECIOUNITARIO', 'PRECIOUSD'},
    'STOCK': {'STOCK', 'CANTIDAD', 'EXISTENCIA', 'INVENTARIO'},
}


def _norm(texto):
    s = unicodedata.normalize('NFKD', str(texto))
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r'[^A-Z0-9]', '', s.upper())


def _es_columna_imagen(n):
    return 'IMAGEN' in n or 'FOTO' in n or ('URL' in n and 'SKU' not in n)


def _recuperar_hipervinculos(archivo, df, nombre_col):
    """Si la celda tiene un hipervínculo con texto visible, recupera la URL real."""
    try:
        from openpyxl import load_workbook
        archivo.seek(0)
        ws = load_workbook(archivo).active
        idx = list(df.columns).index(nombre_col)
        for i, fila in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row)):
            if i >= len(df) or idx >= len(fila):
                break
            celda = fila[idx]
            if celda.hyperlink and celda.hyperlink.target:
                df.at[df.index[i], nombre_col] = celda.hyperlink.target
    except Exception:
        pass  # opcional: si falla, seguimos con lo que leyó pandas
    return df


def cargar_inventario_excel(ruta_archivo):
    if ruta_archivo is None:
        return None

    try:
        df = pd.read_excel(ruta_archivo)
    except Exception as e:
        st.error(f"No se pudo leer el archivo Excel: {e}")
        return None

    # Mapeo flexible de encabezados -> nombres estándar
    renombres, usados, col_img_original = {}, set(), None
    for col in df.columns:
        n = _norm(col)
        asignada = False
        for canonico, alias in ALIAS.items():
            if canonico not in usados and n in alias:
                renombres[col] = canonico
                usados.add(canonico)
                asignada = True
                break
        if not asignada and 'IMAGEN' not in usados and _es_columna_imagen(n):
            renombres[col] = 'IMAGEN'
            usados.add('IMAGEN')
            col_img_original = col

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in usados]
    if faltantes:
        st.error(f"Faltan columnas: {faltantes}. Columnas detectadas en tu archivo: {list(df.columns)}")
        return None

    if col_img_original is not None:
        df = _recuperar_hipervinculos(ruta_archivo, df, col_img_original)

    df = df.rename(columns=renombres)
    columnas = COLUMNAS_REQUERIDAS + (['IMAGEN'] if 'IMAGEN' in df.columns else [])
    df_final = df[columnas].copy()

    if 'IMAGEN' not in df_final.columns:
        df_final['IMAGEN'] = None
    df_final['IMAGEN'] = df_final['IMAGEN'].apply(normalizar_url_imagen)

    df_final['SKU'] = df_final['SKU'].astype(str).str.strip().str.upper()
    df_final['PRECIO_USD'] = pd.to_numeric(df_final['PRECIO_USD'], errors='coerce')
    df_final['STOCK'] = pd.to_numeric(df_final['STOCK'], errors='coerce').fillna(0)

    # Descarta filas vacías o sin precio válido (PRECIO_USD es NOT NULL en la BD)
    df_final = df_final[(df_final['SKU'] != '') & (df_final['SKU'].str.lower() != 'nan')]
    df_final = df_final.dropna(subset=['PRECIO_USD']).reset_index(drop=True)

    return df_final

# import streamlit as st
# import pandas as pd
# from modulos.imagenes import normalizar_url_imagen   # arriba del archivo

# def cargar_inventario_excel(ruta_archivo):
#     # --- LÓGICA DE PROCESAMIENTO DE EXCEL/CSV SUBIDO POR EL PROVEEDOR ---
#     if ruta_archivo is not None:
#         try:
#             # 1. Leer el archivo Excel
#             df_temp = pd.read_excel(ruta_archivo)
            
#             # 2. Normalizar nombres de columnas a mayúsculas y quitar espacios extra
#             df_temp.columns = [str(col).strip().upper() for col in df_temp.columns]
            
#             # 3. Definir las columnas obligatorias mínimo requeridas
#             columnas_requeridas = ['SKU', 'MATERIAL', 'PRECIO_USD', 'STOCK']
            
#             # Validar que al menos las columnas obligatorias estén presentes
#             if all(col in df_temp.columns for col in columnas_requeridas):
                
#                 # 4. Definir qué columnas conservar
#                 columnas_a_conservar = columnas_requeridas.copy()
                
#                 col_img_encontrada = None
#                 for col in df_temp.columns:
#                     col_lower=col.lower()
#                     if 'imagen' in col_lower or 'foto' in col_lower or ('url' in col_lower and 'sku' not in col_lower):
#                         col_img_encontrada = col
#                         break
                
#                 # Si encontró la columna, la renombramos a la estándar del sistema
#                 if col_img_encontrada:
#                     df_temp.rename(columns={col_img_encontrada: 'IMAGEN'}, inplace=True)
#                     columnas_a_conservar.append('IMAGEN')
#                 # Extraemos las columnas válidas
#                 df_final = df_temp[columnas_a_conservar].copy()
                
#                 # Si 'IMAGEN' no venía en el Excel, la creamos vacía para no romper el DataFrame maestro
#                 if 'IMAGEN' not in df_final.columns:
#                     df_final['IMAGEN'] = None
#                 df_final['IMAGEN'] = df_final['IMAGEN'].apply(normalizar_url_imagen)   
                
#                 df_final['SKU'] = df_final['SKU'].astype(str).str.strip().str.upper()
#                 df_final['PRECIO_USD'] = pd.to_numeric(df_final['PRECIO_USD'], errors='coerce')
#                 df_final['STOCK'] = pd.to_numeric(df_final['STOCK'], errors='coerce').fillna(0)

#                 # Descarta filas sin SKU o sin precio válido (PRECIO_USD es NOT NULL en la BD)
#                 df_final = df_final[(df_final['SKU'] != '') & (df_final['SKU'].str.lower() != 'nan')]
#                 df_final = df_final.dropna(subset=['PRECIO_USD']).reset_index(drop=True) 
                
#                 return df_final
                
#             else:
#                 st.error("❌ El archivo Excel no contiene todas las columnas requeridas (SKU, MATERIAL, PRECIO_USD, STOCK).")
#                 return None
            
#         except FileNotFoundError:
#             st.error(f"[Error] No se encontró el archivo en la ruta: {ruta_archivo}")
#             return None
#         except ValueError as error_valor:
#             st.error(f"[Error de Formato] El Excel no tiene la estructura correcta. Detalles: {error_valor}")
#             return None
#         except Exception as e:
#             st.error(f"Error al procesar el archivo Excel: {e}")
#             return None
        
    
    
    # """
    # Lee un archivo Excel y verifica que tenga la estructura correcta.
    # Retorna un DataFrame de Pandas o None si hay un error.
    # """
    # columnas_requeridas = ['SKU', 'Material', 'Precio_USD', 'Stock']
    
    # try:
    #     # Intentamos leer el archivo Excel
    #     df = pd.read_excel(ruta_archivo)
        
    #     # Limpiamos los nombres de las columnas (quitamos espacios y ponemos en mayúsculas)
    #     df.columns = df.columns.str.strip().str.upper()
    #     columnas_requeridas_upper = [col.upper() for col in columnas_requeridas]
        
    #     # Validamos que todas las columnas necesarias existan
    #     for col in columnas_requeridas_upper:
    #         if col not in df.columns:
    #             # Si falta una columna, forzamos un error a propósito
    #             raise ValueError(f"Falta la columna obligatoria: {col}")
        
    #     # Si todo está bien, aseguramos que los números sean números y no texto
    #     df['PRECIO_USD'] = pd.to_numeric(df['PRECIO_USD'], errors='coerce')
    #     df['STOCK'] = pd.to_numeric(df['STOCK'], errors='coerce').fillna(0)
        
    #     return df

    # except FileNotFoundError:
    #     print(f"[Error] No se encontró el archivo en la ruta: {ruta_archivo}")
    #     return None
    # except ValueError as error_valor:
    #     print(f"[Error de Formato] El Excel no tiene la estructura correcta. Detalles: {error_valor}")
    #     return None
    # except Exception as e:
    #     print(f"[Error Inesperado] Ocurrió un problema al leer el Excel: {e}")
    #     return None 
    
